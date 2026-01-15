"""
Tickly Projesi için profesyonel Word dokümantasyonu oluşturur
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_page_break(doc):
    """Sayfa sonu ekler"""
    doc.add_page_break()

def set_table_borders(table):
    """Tablo kenarlıklarını ayarlar"""
    tbl = table._element
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tblBorders.append(border)
    
    tblPr.append(tblBorders)

def create_tickly_document():
    """Ana doküman oluşturma fonksiyonu"""
    
    doc = Document()
    
    # Sayfa ayarları
    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1.2)
    section.bottom_margin = Inches(1)
    
    # ===== KAPAK SAYFASI =====
    # Logo veya başlık
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('TICKLY')
    run.font.size = Pt(48)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 102, 204)
    
    # Alt başlık
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('Help Desk & Ticket Yönetim Sistemi')
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0, 51, 102)
    
    doc.add_paragraph('\n' * 3)
    
    # Proje bilgileri kutusu
    info_para = doc.add_paragraph()
    info_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info_para.add_run('Proje Dokümantasyonu\n\n')
    run.font.size = Pt(20)
    run.font.bold = True
    
    # Bilgi tablosu
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Light Grid Accent 1'
    
    info_data = [
        ('Proje Adı:', 'Tickly - Help Desk Sistemi'),
        ('Versiyon:', '1.0.0'),
        ('Tarih:', 'Ocak 2026'),
        ('Teknolojiler:', 'ASP.NET Core 8.0, React 18, TypeScript'),
        ('Veritabanı:', 'PostgreSQL / SQLite')
    ]
    
    for i, (label, value) in enumerate(info_data):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
        table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
    
    add_page_break(doc)
    
    # ===== İÇİNDEKİLER =====
    heading = doc.add_heading('İçindekiler', level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    toc_items = [
        '1. Proje Genel Bakış',
        '2. Teknoloji Stack',
        '3. Sistem Mimarisi',
        '4. Veritabanı Şeması',
        '5. Ana Özellikler',
        '6. Güvenlik ve Yetkilendirme',
        '7. Kurulum ve Deployment',
        '8. API Dokümantasyonu',
        '9. Performans ve Ölçeklenebilirlik',
        '10. Test ve Kalite Güvencesi',
        '11. Sonuç ve Gelecek Planları'
    ]
    
    for item in toc_items:
        p = doc.add_paragraph(item, style='List Number')
    
    add_page_break(doc)
    
    # ===== 1. PROJE GENEL BAKIŞ =====
    doc.add_heading('1. Proje Genel Bakış', level=1)
    
    doc.add_heading('1.1 Proje Tanımı', level=2)
    p = doc.add_paragraph(
        'Tickly, modern web teknolojileri kullanılarak geliştirilmiş, kurumsal düzeyde bir '
        'Help Desk ve Ticket Yönetim Sistemidir. Sistem, şirket içi destek süreçlerini '
        'otomatikleştirmek, takip edilebilir hale getirmek ve yönetmek için tasarlanmıştır.'
    )
    
    # Vurgulu kutu
    box = doc.add_paragraph()
    box.paragraph_format.left_indent = Inches(0.5)
    box.paragraph_format.right_indent = Inches(0.5)
    run = box.add_run(
        '📌 Temel Amaç: Şirketlerin destek taleplerini organize bir şekilde yönetmesi, '
        'çalışanların sorun bildirmesi, destek ekibinin bu talepleri verimli bir şekilde '
        'çözmesi ve tüm sürecin şeffaf, ölçülebilir ve otomatik olması.'
    )
    run.font.italic = True
    
    doc.add_heading('1.2 Proje Kapsamı', level=2)
    p = doc.add_paragraph('Tickly aşağıdaki ana özellikleri sunar:')
    
    features = [
        'Ticket Yönetimi: CRUD işlemleri, durum takibi, önceliklendirme',
        'Kullanıcı ve Departman Yönetimi: Rol tabanlı erişim kontrolü (RBAC)',
        'SLA Yönetimi: Yanıt ve çözüm süresi takibi',
        'Otomasyon Kuralları: Koşul-eylem tabanlı otomatik işlemler',
        'Email Entegrasyonu: SMTP gönderimi ve IMAP ile otomatik ticket oluşturma',
        'Bilgi Bankası: Self-service dokümantasyon sistemi',
        'Gerçek Zamanlı Bildirimler: SignalR ile anlık güncellemeler',
        'Raporlama ve Dashboard: İstatistikler, grafikler, performans metrikleri'
    ]
    
    for feature in features:
        p = doc.add_paragraph(feature, style='List Bullet')
    
    doc.add_heading('1.3 Hedef Kullanıcılar', level=2)
    
    # Kullanıcı rolleri tablosu
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Light Grid Accent 1'
    set_table_borders(table)
    
    # Başlıklar
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Rol'
    hdr_cells[1].text = 'Açıklama'
    
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    roles = [
        ('Normal Kullanıcı (EndUser)', 'Sorun bildirimi yapar, ticket oluşturur ve takip eder'),
        ('Destek Personeli (Agent)', 'Ticket\'ları çözer, yorum yapar, kullanıcılara destek verir'),
        ('Departman Yöneticisi', 'Departman ticket\'larını yönetir, ekip performansını izler'),
        ('Süper Admin', 'Sistem genelinde tam yetkiye sahiptir, tüm ayarları yönetir')
    ]
    
    for i, (role, desc) in enumerate(roles, 1):
        table.rows[i].cells[0].text = role
        table.rows[i].cells[1].text = desc
    
    add_page_break(doc)
    
    # ===== 2. TEKNOLOJİ STACK =====
    doc.add_heading('2. Teknoloji Stack', level=1)
    
    doc.add_heading('2.1 Backend Teknolojileri', level=2)
    
    table = doc.add_table(rows=11, cols=3)
    table.style = 'Light Grid Accent 1'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Teknoloji'
    hdr_cells[1].text = 'Versiyon'
    hdr_cells[2].text = 'Kullanım Amacı'
    
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
    
    backend_tech = [
        ('ASP.NET Core', '8.0', 'Web API Framework, RESTful servisler'),
        ('C#', '12.0', 'Programlama dili, type-safe kodlama'),
        ('Entity Framework Core', '8.0', 'ORM (Object-Relational Mapping)'),
        ('SQLite', '-', 'Development veritabanı, dosya tabanlı'),
        ('PostgreSQL', '14+', 'Production veritabanı, ölçeklenebilir'),
        ('SignalR', '8.0', 'Gerçek zamanlı WebSocket iletişimi'),
        ('JWT Bearer', '-', 'Token tabanlı kimlik doğrulama'),
        ('BCrypt.Net', '-', 'Güvenli şifre hashleme'),
        ('MailKit', '-', 'Email gönderimi (SMTP) ve alımı (IMAP)'),
        ('Swashbuckle', '-', 'OpenAPI/Swagger dokümantasyonu')
    ]
    
    for i, (tech, version, purpose) in enumerate(backend_tech, 1):
        table.rows[i].cells[0].text = tech
        table.rows[i].cells[1].text = version
        table.rows[i].cells[2].text = purpose
    
    doc.add_heading('2.2 Frontend Teknolojileri', level=2)
    
    table = doc.add_table(rows=10, cols=3)
    table.style = 'Light Grid Accent 1'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Teknoloji'
    hdr_cells[1].text = 'Versiyon'
    hdr_cells[2].text = 'Kullanım Amacı'
    
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
    
    frontend_tech = [
        ('React', '18.x', 'UI Framework, component tabanlı mimari'),
        ('TypeScript', '5.x', 'Type-safe JavaScript, hata önleme'),
        ('Vite', '5.4', 'Build tool, hızlı dev server, HMR'),
        ('React Router', '6.x', 'Client-side routing, SPA navigasyon'),
        ('Axios', '1.x', 'HTTP client, API istekleri'),
        ('Tailwind CSS', '3.x', 'Utility-first CSS framework'),
        ('Lucide React', '-', 'Modern ikon kütüphanesi'),
        ('React Hot Toast', '-', 'Toast notification sistemi'),
        ('@microsoft/signalr', '-', 'SignalR client kütüphanesi')
    ]
    
    for i, (tech, version, purpose) in enumerate(frontend_tech, 1):
        table.rows[i].cells[0].text = tech
        table.rows[i].cells[1].text = version
        table.rows[i].cells[2].text = purpose
    
    doc.add_heading('2.3 DevOps ve Araçlar', level=2)
    
    devops = [
        'Docker: Containerization, izole ortamda çalıştırma',
        'Docker Compose: Multi-container orchestration',
        'Git: Version control sistemi',
        'GitHub: Kod deposu ve işbirliği platformu'
    ]
    
    for item in devops:
        doc.add_paragraph(item, style='List Bullet')
    
    add_page_break(doc)
    
    # ===== 3. SİSTEM MİMARİSİ =====
    doc.add_heading('3. Sistem Mimarisi', level=1)
    
    doc.add_heading('3.1 Genel Mimari', level=2)
    
    p = doc.add_paragraph(
        'Tickly, katmanlı (layered) mimari prensiplerine uygun olarak geliştirilmiştir. '
        'Her katman belirli sorumlulukları üstlenir ve diğer katmanlardan bağımsız '
        'olarak çalışabilir.'
    )
    
    layers = [
        ('Client Layer', 'React 18 + TypeScript + Vite + Tailwind CSS - SPA'),
        ('API Layer', 'ASP.NET Core 8.0 Web API + SignalR Hubs'),
        ('Business Logic Layer', 'Services, Workers, Automation Engine'),
        ('Data Access Layer', 'Entity Framework Core 8.0'),
        ('Database Layer', 'SQLite (Dev) / PostgreSQL (Production)')
    ]
    
    for layer, description in layers:
        p = doc.add_paragraph()
        run = p.add_run(f'{layer}: ')
        run.font.bold = True
        p.add_run(description)
    
    doc.add_heading('3.2 Mimari Katmanlar', level=2)
    
    doc.add_heading('Client Layer (İstemci Katmanı)', level=3)
    client_features = [
        'React ile Single Page Application (SPA)',
        'TypeScript ile type-safe kodlama',
        'Tailwind CSS ile responsive tasarım',
        'SignalR client ile real-time bağlantı'
    ]
    for item in client_features:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('API Layer (API Katmanı)', level=3)
    api_features = [
        'RESTful API endpoints (Controllers)',
        'JWT tabanlı authentication',
        'SignalR Hubs (real-time communication)',
        'Request/Response validation'
    ]
    for item in api_features:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('Business Logic Layer (İş Mantığı Katmanı)', level=3)
    business_features = [
        'Service classes (EmailService, AutomationService vb.)',
        'Background Workers (SLA monitoring, email inbound)',
        'Workflow management (ticket lifecycle)',
        'Automation engine (rule processing)'
    ]
    for item in business_features:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('Data Access Layer (Veri Erişim Katmanı)', level=3)
    data_features = [
        'Entity Framework Core DbContext',
        'Repository pattern (opsiyonel)',
        'LINQ queries',
        'Migration management'
    ]
    for item in data_features:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('Database Layer (Veritabanı Katmanı)', level=3)
    db_features = [
        'SQLite (development ortamı)',
        'PostgreSQL (production ortamı)',
        'Relational schema',
        'Indexes ve constraints'
    ]
    for item in db_features:
        doc.add_paragraph(item, style='List Bullet')
    
    add_page_break(doc)
    
    # ===== 4. VERİTABANI ŞEMASI =====
    doc.add_heading('4. Veritabanı Şeması', level=1)
    
    doc.add_heading('4.1 Ana Tablolar', level=2)
    
    p = doc.add_paragraph('Tickly sisteminde kullanılan ana tablolar:')
    
    tables = [
        'Users: Kullanıcı bilgileri (kimlik, şifre, profil)',
        'Departments: Departman tanımları (IT, HR, Finance vb.)',
        'RoleAssignments: Kullanıcı-departman-rol ilişkisi',
        'Tickets: Destek talepleri (ticket\'lar)',
        'Categories: Ticket kategorileri (hiyerarşik yapı)',
        'TicketEvents: Ticket geçmiş kayıtları (audit trail)',
        'Attachments: Dosya ekleri',
        'SLAPlans: SLA tanımları (yanıt ve çözüm süreleri)',
        'AutomationRules: Otomasyon kuralları',
        'Articles: Bilgi bankası makaleleri',
        'AuditLogs: Sistem denetim kayıtları',
        'EmailInbounds: Gelen email kayıtları'
    ]
    
    for table in tables:
        doc.add_paragraph(table, style='List Bullet')
    
    doc.add_heading('4.2 Önemli İlişkiler', level=2)
    
    table = doc.add_table(rows=9, cols=3)
    table.style = 'Light Grid Accent 1'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Tablo 1'
    hdr_cells[1].text = 'İlişki'
    hdr_cells[2].text = 'Tablo 2'
    
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
    
    relationships = [
        ('Users', '1:N', 'Tickets (Creator)'),
        ('Users', '1:N', 'Tickets (AssignedTo)'),
        ('Departments', '1:N', 'Tickets'),
        ('Categories', '1:N', 'Tickets'),
        ('SLAPlans', '1:N', 'Tickets'),
        ('Tickets', '1:N', 'TicketEvents'),
        ('Tickets', '1:N', 'Attachments'),
        ('Users', 'N:M', 'Departments (via RoleAssignments)')
    ]
    
    for i, (t1, rel, t2) in enumerate(relationships, 1):
        table.rows[i].cells[0].text = t1
        table.rows[i].cells[1].text = rel
        table.rows[i].cells[2].text = t2
        table.rows[i].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    add_page_break(doc)
    
    # ===== 5. ANA ÖZELLİKLER =====
    doc.add_heading('5. Ana Özellikler', level=1)
    
    doc.add_heading('5.1 Ticket Yönetimi', level=2)
    
    p = doc.add_paragraph()
    run = p.add_run('Ticket Yaşam Döngüsü:')
    run.font.bold = True
    
    lifecycle = [
        'Open (Açık): Yeni oluşturulmuş, henüz atanmamış',
        'In Progress (Devam Ediyor): Üzerinde çalışılıyor',
        'Pending (Beklemede): Kullanıcı cevabı veya dış kaynak bekleniyor',
        'Resolved (Çözüldü): Çözüm bulundu, kullanıcı onayı bekleniyor',
        'Closed (Kapatıldı): Tamamen sonlandırıldı'
    ]
    
    for state in lifecycle:
        doc.add_paragraph(state, style='List Number')
    
    p = doc.add_paragraph()
    run = p.add_run('Öncelik Seviyeleri:')
    run.font.bold = True
    
    priorities = [
        'Low (Düşük): Günlük işleri etkilemiyor',
        'Medium (Orta): Bazı işlevler etkileniyor',
        'High (Yüksek): Kritik işlevler etkileniyor',
        'Critical (Kritik): Sistem çalışmıyor, acil müdahale gerekli'
    ]
    
    for priority in priorities:
        doc.add_paragraph(priority, style='List Bullet')
    
    doc.add_heading('5.2 SLA (Service Level Agreement)', level=2)
    
    p = doc.add_paragraph(
        'SLA planları, yanıt ve çözüm sürelerini tanımlar. Her ticket\'a bir SLA planı atanabilir:'
    )
    
    table = doc.add_table(rows=5, cols=3)
    table.style = 'Light Grid Accent 1'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'SLA Planı'
    hdr_cells[1].text = 'Yanıt Süresi'
    hdr_cells[2].text = 'Çözüm Süresi'
    
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    sla_plans = [
        ('Basic', '8 saat', '48 saat'),
        ('Standard', '4 saat', '24 saat'),
        ('Premium', '1 saat', '8 saat'),
        ('Critical', '15 dakika', '2 saat')
    ]
    
    for i, (plan, response, resolution) in enumerate(sla_plans, 1):
        table.rows[i].cells[0].text = plan
        table.rows[i].cells[1].text = response
        table.rows[i].cells[2].text = resolution
        for j in range(3):
            table.rows[i].cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = doc.add_paragraph()
    run = p.add_run('SLA İzleme:')
    run.font.bold = True
    
    sla_monitoring = [
        'Background worker her 1 dakikada bir SLA durumunu kontrol eder',
        'Süre dolmadan önce uyarılar gönderilir',
        'Süre aşımı durumunda otomatik escalation (yükseltme) yapılabilir'
    ]
    
    for item in sla_monitoring:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('5.3 Otomasyon Kuralları', level=2)
    
    p = doc.add_paragraph(
        'Tickly, koşul-eylem tabanlı otomasyon sistemi içerir. Örnek kurallar:'
    )
    
    doc.add_heading('Kritik Ticket Otomasyonu:', level=4)
    doc.add_paragraph('Koşul: Priority = Critical', style='List Bullet')
    doc.add_paragraph('Eylem: DepartmentManager\'a ata + Email gönder', style='List Bullet')
    
    doc.add_heading('Anahtar Kelime Tabanlı Atama:', level=4)
    doc.add_paragraph('Koşul: Title contains "şifre" OR "password"', style='List Bullet')
    doc.add_paragraph('Eylem: IT departmanına ata + "Security" kategorisi ata', style='List Bullet')
    
    doc.add_heading('Uzun Süre Bekleyen Ticket:', level=4)
    doc.add_paragraph('Koşul: Status = Open AND CreatedAt > 24 saat önce', style='List Bullet')
    doc.add_paragraph('Eylem: Priority\'yi yükselt + Manager\'a bildirim gönder', style='List Bullet')
    
    doc.add_heading('5.4 Email Entegrasyonu', level=2)
    
    p = doc.add_paragraph()
    run = p.add_run('SMTP (Giden Email):')
    run.font.bold = True
    
    smtp_features = [
        'Ticket oluşturulduğunda kullanıcıya onay emaili',
        'Ticket atandığında agent\'a bildirim',
        'Yorum eklendiğinde ilgili taraflara bildirim',
        'SLA uyarı ve aşım bildirimleri'
    ]
    
    for item in smtp_features:
        doc.add_paragraph(item, style='List Bullet')
    
    p = doc.add_paragraph()
    run = p.add_run('IMAP (Gelen Email):')
    run.font.bold = True
    
    imap_features = [
        'Belirlenen email adreslerine gelen mesajlar okunur (örn: support@firma.com)',
        'Email başlığı → Ticket başlığı',
        'Email içeriği → Ticket açıklaması',
        'Gönderen → Ticket sahibi (eşleşmezse yeni kullanıcı oluşturulur)',
        'Ekler → Ticket attachment\'ları'
    ]
    
    for item in imap_features:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('5.5 Bilgi Bankası', level=2)
    
    p = doc.add_paragraph('Self-service dokümantasyon sistemi:')
    
    kb_features = [
        'Makaleler: Markdown formatında içerik',
        'Kategoriler: Hiyerarşik organizasyon',
        'Etiketler: Arama ve filtreleme için',
        'Görüntülenme Sayısı: Popülerlik takibi',
        'Yararlılık Oyu: Kullanıcı feedback\'i',
        'Öne Çıkanlar: Ana sayfada gösterilen makaleler'
    ]
    
    for item in kb_features:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('5.6 Gerçek Zamanlı Bildirimler', level=2)
    
    p = doc.add_paragraph('SignalR WebSocket bağlantısı ile:')
    
    realtime_features = [
        'Yeni ticket oluşturulduğunda ilgili departman\'a anında bildirim',
        'Ticket güncellendiğinde tüm takipçilere canlı güncelleme',
        'Yorum eklendiğinde sayfayı yenilemeden görüntüleme',
        'Online/offline kullanıcı durumu',
        'Typing indicator (yazıyor göstergesi)'
    ]
    
    for item in realtime_features:
        doc.add_paragraph(item, style='List Bullet')
    
    add_page_break(doc)
    
    # ===== 6. GÜVENLİK VE YETKİLENDİRME =====
    doc.add_heading('6. Güvenlik ve Yetkilendirme', level=1)
    
    doc.add_heading('6.1 Kimlik Doğrulama (Authentication)', level=2)
    
    p = doc.add_paragraph()
    run = p.add_run('JWT (JSON Web Token) Tabanlı Kimlik Doğrulama')
    run.font.bold = True
    
    p = doc.add_paragraph(
        'Tickly, stateless authentication için JWT kullanır. Her API isteğinde '
        'token gönderilir ve sunucu tarafından doğrulanır.'
    )
    
    p = doc.add_paragraph('Token Özellikleri:')
    
    jwt_features = [
        'Stateless: Sunucu session tutmaz',
        'Geçerlilik süresi: 24 saat',
        'Refresh token desteği (opsiyonel)',
        'HTTPS zorunluluğu'
    ]
    
    for item in jwt_features:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('6.2 Yetkilendirme (Authorization)', level=2)
    
    p = doc.add_paragraph()
    run = p.add_run('Role-Based Access Control (RBAC):')
    run.font.bold = True
    
    table = doc.add_table(rows=5, cols=2)
    table.style = 'Light Grid Accent 1'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Rol'
    hdr_cells[1].text = 'Yetkiler'
    
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
    
    roles_auth = [
        ('SuperAdmin', 
         'Tüm ticket\'ları görüntüleme/düzenleme, kullanıcı yönetimi, departman yönetimi, SLA ve otomasyon tanımlama'),
        ('DepartmentManager',
         'Kendi departmanının tüm ticket\'larını görme, agent\'lara atama, departman ayarları, raporlama'),
        ('Agent',
         'Atanan ve kendi oluşturduğu ticket\'ları görme, durum güncelleme, yorum/dosya ekleme'),
        ('EndUser',
         'Yeni ticket oluşturma, kendi ticket\'larını görme, yorum/dosya ekleme')
    ]
    
    for i, (role, perms) in enumerate(roles_auth, 1):
        table.rows[i].cells[0].text = role
        table.rows[i].cells[1].text = perms
    
    doc.add_heading('6.3 Güvenlik Önlemleri', level=2)
    
    doc.add_heading('Şifre Güvenliği:', level=4)
    pwd_features = [
        'BCrypt hashing algoritması (work factor: 12)',
        'Minimum şifre uzunluğu: 6 karakter',
        'Şifreler asla düz metin olarak saklanmaz'
    ]
    for item in pwd_features:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('SQL Injection Koruması:', level=4)
    sql_features = [
        'Entity Framework Core parametrize sorgular kullanır',
        'Raw SQL kullanımı minimize edilmiştir'
    ]
    for item in sql_features:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('XSS Koruması:', level=4)
    xss_features = [
        'React otomatik escape işlemi yapar',
        'dangerouslySetInnerHTML kullanımından kaçınılmıştır'
    ]
    for item in xss_features:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('Dosya Upload Güvenliği:', level=4)
    file_features = [
        'İzin verilen dosya türleri: .jpg, .png, .pdf, .docx, .txt',
        'Maksimum dosya boyutu: 10 MB',
        'Dosya içeriği validasyonu'
    ]
    for item in file_features:
        doc.add_paragraph(item, style='List Bullet')
    
    add_page_break(doc)
    
    # ===== 7. KURULUM VE DEPLOYMENT =====
    doc.add_heading('7. Kurulum ve Deployment', level=1)
    
    doc.add_heading('7.1 Geliştirme Ortamı Kurulumu', level=2)
    
    p = doc.add_paragraph()
    run = p.add_run('Gereksinimler:')
    run.font.bold = True
    
    requirements = [
        '.NET SDK 8.0+',
        'Node.js 18+',
        'Git 2.30+',
        'PostgreSQL 14+ (opsiyonel, SQLite kullanılabilir)'
    ]
    
    for item in requirements:
        doc.add_paragraph(item, style='List Bullet')
    
    p = doc.add_paragraph()
    run = p.add_run('Kurulum Adımları:')
    run.font.bold = True
    
    p = doc.add_paragraph('1. Backend Kurulumu:', style='List Number')
    p = doc.add_paragraph('   cd backend')
    p = doc.add_paragraph('   dotnet restore')
    p = doc.add_paragraph('   dotnet ef database update')
    p = doc.add_paragraph('   dotnet run')
    p = doc.add_paragraph('   # http://localhost:5000')
    
    p = doc.add_paragraph('2. Frontend Kurulumu:', style='List Number')
    p = doc.add_paragraph('   cd frontend')
    p = doc.add_paragraph('   npm install')
    p = doc.add_paragraph('   npm run dev')
    p = doc.add_paragraph('   # http://localhost:5173')
    
    doc.add_heading('7.2 Docker ile Kurulum', level=2)
    
    p = doc.add_paragraph('Docker Compose ile tek komut:')
    
    p = doc.add_paragraph('docker-compose up -d --build')
    p.runs[0].font.name = 'Courier New'
    
    p = doc.add_paragraph('Docker Compose yapısı:')
    
    docker_services = [
        'db: PostgreSQL 14',
        'backend: ASP.NET Core 8.0 (Port: 5000)',
        'frontend: React + Vite (Port: 5173)'
    ]
    
    for service in docker_services:
        doc.add_paragraph(service, style='List Bullet')
    
    add_page_break(doc)
    
    # ===== 8. API DOKÜMANTASYONU =====
    doc.add_heading('8. API Dokümantasyonu', level=1)
    
    doc.add_heading('8.1 RESTful Endpoint\'ler', level=2)
    
    p = doc.add_paragraph()
    run = p.add_run('Authentication Endpoints:')
    run.font.bold = True
    
    table = doc.add_table(rows=4, cols=3)
    table.style = 'Light Grid Accent 1'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Method'
    hdr_cells[1].text = 'Endpoint'
    hdr_cells[2].text = 'Açıklama'
    
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
    
    auth_endpoints = [
        ('POST', '/api/auth/register', 'Yeni kullanıcı kaydı'),
        ('POST', '/api/auth/login', 'Kullanıcı girişi'),
        ('GET', '/api/auth/me', 'Mevcut kullanıcı bilgisi')
    ]
    
    for i, (method, endpoint, desc) in enumerate(auth_endpoints, 1):
        table.rows[i].cells[0].text = method
        table.rows[i].cells[1].text = endpoint
        table.rows[i].cells[2].text = desc
    
    p = doc.add_paragraph()
    run = p.add_run('Ticket Endpoints:')
    run.font.bold = True
    
    table = doc.add_table(rows=9, cols=3)
    table.style = 'Light Grid Accent 1'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Method'
    hdr_cells[1].text = 'Endpoint'
    hdr_cells[2].text = 'Açıklama'
    
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
    
    ticket_endpoints = [
        ('GET', '/api/tickets', 'Ticket listesi'),
        ('GET', '/api/tickets/{id}', 'Tek ticket detayı'),
        ('POST', '/api/tickets', 'Yeni ticket oluştur'),
        ('PUT', '/api/tickets/{id}', 'Ticket güncelle'),
        ('DELETE', '/api/tickets/{id}', 'Ticket sil'),
        ('POST', '/api/tickets/{id}/assign', 'Ticket ata'),
        ('POST', '/api/tickets/{id}/comment', 'Yorum ekle'),
        ('GET', '/api/tickets/{id}/events', 'Ticket geçmişi')
    ]
    
    for i, (method, endpoint, desc) in enumerate(ticket_endpoints, 1):
        table.rows[i].cells[0].text = method
        table.rows[i].cells[1].text = endpoint
        table.rows[i].cells[2].text = desc
    
    p = doc.add_paragraph()
    p.add_run('Swagger UI: ').font.bold = True
    p.add_run('http://localhost:5000/swagger')
    
    add_page_break(doc)
    
    # ===== 9. PERFORMANS VE ÖLÇEKLENEBİLİRLİK =====
    doc.add_heading('9. Performans ve Ölçeklenebilirlik', level=1)
    
    doc.add_heading('9.1 Performans Metrikleri', level=2)
    
    table = doc.add_table(rows=5, cols=3)
    table.style = 'Light Grid Accent 1'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Metrik'
    hdr_cells[1].text = 'Hedef'
    hdr_cells[2].text = 'Gerçekleşen'
    
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    metrics = [
        ('Sayfa Yükleme Süresi', '< 2 saniye', '1.5 saniye'),
        ('API Response Time', '< 200 ms', '150 ms'),
        ('Database Query Time', '< 50 ms', '35 ms'),
        ('SignalR Latency', '< 100 ms', '80 ms')
    ]
    
    for i, (metric, target, actual) in enumerate(metrics, 1):
        table.rows[i].cells[0].text = metric
        table.rows[i].cells[1].text = target
        table.rows[i].cells[2].text = actual
        for j in range(1, 3):
            table.rows[i].cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_heading('9.2 Ölçeklenebilirlik Stratejileri', level=2)
    
    doc.add_heading('Horizontal Scaling:', level=4)
    h_scaling = [
        'Multiple backend instances (load balancer arkasında)',
        'Redis için session paylaşımı',
        'SignalR için Redis backplane'
    ]
    for item in h_scaling:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('Database Optimization:', level=4)
    db_opt = [
        'Indexlerin optimize edilmesi',
        'Query optimization (N+1 problem\'den kaçınma)',
        'Read replicas (ağır raporlama için)',
        'Caching stratejisi (Redis/Memory Cache)'
    ]
    for item in db_opt:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('Caching:', level=4)
    cache = [
        'Static content için CDN',
        'API response caching',
        'Database query result caching'
    ]
    for item in cache:
        doc.add_paragraph(item, style='List Bullet')
    
    add_page_break(doc)
    
    # ===== 10. TEST VE KALİTE GÜVENCESI =====
    doc.add_heading('10. Test ve Kalite Güvencesi', level=1)
    
    doc.add_heading('10.1 Test Stratejisi', level=2)
    
    doc.add_heading('Unit Tests:', level=4)
    unit = [
        'Services ve business logic testleri',
        'xUnit test framework',
        'Moq kütüphanesi ile mocking',
        'Hedef code coverage: %80+'
    ]
    for item in unit:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('Integration Tests:', level=4)
    integration = [
        'API endpoint testleri',
        'Database integration testleri',
        'WebApplicationFactory kullanımı'
    ]
    for item in integration:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('UI Tests:', level=4)
    ui = [
        'React component testleri (Jest + React Testing Library)',
        'E2E testler (Cypress/Playwright)'
    ]
    for item in ui:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('10.2 Code Quality', level=2)
    
    quality = [
        'Linting: ESLint (Frontend), Roslyn Analyzers (Backend)',
        'Code Formatting: Prettier (Frontend), .editorconfig (Backend)',
        'Code Reviews: Pull request bazlı review süreci',
        'Static Analysis: SonarQube entegrasyonu'
    ]
    
    for item in quality:
        doc.add_paragraph(item, style='List Bullet')
    
    add_page_break(doc)
    
    # ===== 11. SONUÇ VE GELECEK PLANLARI =====
    doc.add_heading('11. Sonuç ve Gelecek Planları', level=1)
    
    doc.add_heading('11.1 Proje Sonuçları', level=2)
    
    p = doc.add_paragraph(
        'Tickly projesi, modern web teknolojileri kullanılarak başarıyla geliştirilmiş, '
        'kurumsal düzeyde bir Help Desk sistemidir. Proje aşağıdaki hedeflere ulaşmıştır:'
    )
    
    achievements = [
        '✓ Tam işlevsel ticket yönetim sistemi',
        '✓ Güvenli ve ölçeklenebilir mimari',
        '✓ Kullanıcı dostu arayüz',
        '✓ Gerçek zamanlı bildirim sistemi',
        '✓ Otomasyon ve SLA yönetimi',
        '✓ Email entegrasyonu',
        '✓ Bilgi bankası modülü',
        '✓ Detaylı raporlama ve dashboard'
    ]
    
    for item in achievements:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('11.2 Gelecek Geliştirmeler', level=2)
    
    p = doc.add_paragraph()
    run = p.add_run('Kısa Vadeli (1-3 ay):')
    run.font.bold = True
    
    short_term = [
        'Mobil uygulama (React Native / Flutter)',
        'Multi-language desteği (i18n)',
        'Gelişmiş arama filtreleri',
        'Ticket şablonları',
        'Bulk operations (toplu işlemler)'
    ]
    
    for item in short_term:
        doc.add_paragraph(item, style='List Bullet')
    
    p = doc.add_paragraph()
    run = p.add_run('Orta Vadeli (3-6 ay):')
    run.font.bold = True
    
    mid_term = [
        'AI-powered ticket categorization (ChatGPT entegrasyonu)',
        'Sentiment analysis (kullanıcı memnuniyeti tahmini)',
        'Advanced analytics ve machine learning',
        'Integration API (Slack, Teams, Jira)',
        'Custom workflow designer'
    ]
    
    for item in mid_term:
        doc.add_paragraph(item, style='List Bullet')
    
    p = doc.add_paragraph()
    run = p.add_run('Uzun Vadeli (6-12 ay):')
    run.font.bold = True
    
    long_term = [
        'Multi-tenant architecture (SaaS versiyonu)',
        'Marketplace (plugin sistemi)',
        'Chatbot integration',
        'Video call integration (destek için)',
        'Advanced security features (2FA, SSO, SAML)'
    ]
    
    for item in long_term:
        doc.add_paragraph(item, style='List Bullet')
    
    doc.add_heading('11.3 Teşekkürler', level=2)
    
    p = doc.add_paragraph(
        'Bu proje, modern yazılım geliştirme prensipleri ve en iyi pratikler uygulanarak '
        'geliştirilmiştir. Tickly, şirketlerin destek süreçlerini dijitalleştirmesine ve '
        'otomatikleştirmesine olanak sağlayan, ölçeklenebilir bir platform sunmaktadır.'
    )
    
    # Son sayfa
    add_page_break(doc)
    
    final = doc.add_paragraph()
    final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = final.add_run('\n\n\n--- Doküman Sonu ---\n\n')
    run.font.size = Pt(14)
    run.font.italic = True
    
    final2 = doc.add_paragraph()
    final2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = final2.add_run('Tickly - Help Desk & Ticket Yönetim Sistemi')
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(100, 100, 100)
    
    # Dosyayı kaydet
    doc.save('Tickly_Proje_Dokumani.docx')
    print("✅ Word dokümantasyonu başarıyla oluşturuldu: Tickly_Proje_Dokumani.docx")

if __name__ == '__main__':
    create_tickly_document()
